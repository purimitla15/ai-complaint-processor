from pathlib import Path

import pytest
from docx import Document

from src.ingestion.document_loader import DocumentLoadError, discover_documents, load_document

SAMPLE_TEXT = "My name is Test User and my order ORD-1 arrived damaged. Please help."


def test_loads_txt(tmp_path: Path):
    path = tmp_path / "case.txt"
    path.write_text(SAMPLE_TEXT, encoding="utf-8")
    document = load_document(path)
    assert document.text == SAMPLE_TEXT
    assert document.file_type == "txt"
    assert document.document_id == "case"


def test_loads_docx_including_tables(tmp_path: Path):
    path = tmp_path / "case.docx"
    doc = Document()
    doc.add_paragraph(SAMPLE_TEXT)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Order ID", "ORD-1"
    doc.save(str(path))

    document = load_document(path)
    assert SAMPLE_TEXT in document.text
    assert "Order ID | ORD-1" in document.text


def test_corrupt_pdf_raises_load_error(tmp_path: Path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.4 truncated")
    with pytest.raises(DocumentLoadError):
        load_document(path)


def test_empty_file_raises_load_error(tmp_path: Path):
    path = tmp_path / "empty.txt"
    path.write_text("   ", encoding="utf-8")
    with pytest.raises(DocumentLoadError, match="no usable text"):
        load_document(path)


def test_discover_separates_unsupported_files(tmp_path: Path):
    (tmp_path / "a.txt").write_text(SAMPLE_TEXT)
    (tmp_path / "b.xlsx").write_bytes(b"x")
    (tmp_path / ".hidden.txt").write_text(SAMPLE_TEXT)

    supported, skipped = discover_documents(tmp_path)
    assert [p.name for p in supported] == ["a.txt"]
    assert [p.name for p in skipped] == ["b.xlsx"]


def test_missing_data_dir_raises(tmp_path: Path):
    with pytest.raises(DocumentLoadError, match="not found"):
        discover_documents(tmp_path / "does-not-exist")
