"""Generate the sample documents in data/.

All customers, companies, emails and phone numbers are fictional.
Requires the dev dependencies:  pip install -r requirements-dev.txt

Run from the project root:  python scripts/generate_sample_data.py
"""

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

DATA_DIR = Path(__file__).resolve().parent.parent / "data"

COMPLAINT_001_PDF = """\
To: Customer Care, NovaHome Electronics
Date: 02 September 2026
Subject: Defective washing machine - demand for replacement

I am Ravi Kumar, residing at 14, 5th Cross, Jayanagar, Bengaluru. My email is
ravi.kumar@example.com and my mobile number is +91 98450 12345.

I purchased a NovaHome AquaClean 7kg front-load washing machine (Order No. NH-558120)
on 03 August 2026. The machine stopped working on 13 August 2026 - the drum does not
spin and it shows error code E21.

Your technician visited on 18 August 2026 and replaced the drain pump, but the same
problem returned within two days. A second visit was scheduled for 25 August but nobody
turned up. I have called your helpline three times (reference CRM-77231) without any result.

I request a full replacement of the machine. If this is not resolved within 7 days I will
be forced to approach the Consumer Disputes Redressal Commission.

Enclosed: copy of the invoice and the technician job card.

Ravi Kumar"""

COMPLAINT_002_TXT = """\
From: Priya Sharma <priya.sharma@example.com>
Phone: 080-4123-9876
Subject: Double charge on my FiberNet broadband bill - RESOLVED, thank you

Hello Support Team,

I had written last week about being charged twice (Rs. 1,178 x 2) for my FiberNet
Broadband 200 Mbps plan for the month of August 2026, account number FN-204417.

I am writing to confirm that the duplicate amount of Rs. 1,178 has now been credited
back to my card on 09 September 2026, as your agent Mr. Arjun had promised. The issue
is fully resolved from my side.

However, I would like to point out that it took three calls and eight days to get this
fixed. Please improve your billing checks so this does not happen to other customers.

Regards,
Priya Sharma
"""

COMPLAINT_003_DOCX = {
    "title": "Complaint Regarding Delayed Laptop Delivery",
    "paragraphs": [
        "Dear ShopKart Customer Support,",
        "My name is Mohammed Irfan. I am writing about my order for a laptop which has "
        "still not been delivered even though the promised delivery date has passed.",
        "The tracking page has shown 'In transit - Hyderabad hub' for the last six days. "
        "On 10 September 2026 your chat agent told me that a query had been raised with "
        "the courier partner and that I would receive an update within 48 hours. I have "
        "not received any update since.",
        "I need this laptop for my new job which starts on 22 September 2026. Please "
        "deliver it at the earliest or tell me clearly what is happening with my order.",
        "You can reach me at irfan.m@example.com or 99001 22334.",
        "Thank you,",
        "Mohammed Irfan",
    ],
    "table": [
        ("Order ID", "SK-2026-884512"),
        ("Product", "Zenbook Air 14 Laptop (16GB / 512GB)"),
        ("Order date", "01 September 2026"),
        ("Promised delivery", "06 September 2026"),
        ("Amount paid", "Rs. 68,990 (prepaid)"),
    ],
}

COMPLAINT_004_PDF = """\
Support ticket submitted via web form - PayWise Mobile App

Name: Ananya Reddy
Email: ananya.reddy@example.com
Category selected by customer: Account

Message:
Since the app update on 11 September 2026 I am unable to log in to my PayWise account.
After entering the OTP the app shows "Session expired, please try again" and returns to
the login screen. I have reinstalled the app twice and tried on both mobile data and WiFi.

My electricity bill auto-pay is due on 20 September and I cannot check whether it is set
up correctly. I have attached screenshots of the error message.

Please help me regain access to my account."""

COMPLAINT_005_TXT = """\
From: Suresh Iyer <suresh.iyer@example.com>
Subject: Question about extended warranty for my refrigerator

Hi,

I bought a NovaHome FrostFree 340L refrigerator in March 2026 and it is working very well -
very happy with the purchase.

The standard warranty is one year. I wanted to know whether I can buy an extended warranty
now, and if so, what the cost and coverage would be. Does it also cover the compressor?

Thanks,
Suresh Iyer
Mobile: 98860 55443
"""

COMPLAINT_006_DOCX = {
    "title": "URGENT: Refund Not Received - Request Escalation to Manager",
    "paragraphs": [
        "To the Customer Service Manager, StyleStreet Fashion",
        "I am Kavya Nair (kavya.nair@example.com, +91 94470 67890). I returned a pair of "
        "running shoes (Order ST-39021, amount Rs. 4,499) on 12 August 2026 because the "
        "size was incorrect. Your courier picked up the item and the return was marked "
        "'Received at warehouse' on 16 August 2026.",
        "Your policy promises a refund within 7 working days. It has now been more than "
        "four weeks and I have not received my money.",
        "I have followed up three times - by email on 26 August, by phone on 02 September "
        "(ticket SS-10442) and on chat on 10 September. Each time I was told the refund is "
        "'under process', with no timeline.",
        "I am attaching screenshots of the return confirmation and my previous chat "
        "transcripts. I request that this be escalated to a manager immediately and that "
        "the refund be processed without further delay.",
        "Kavya Nair",
    ],
    "table": None,
}


def write_pdf(path: Path, text: str) -> None:
    styles = getSampleStyleSheet()
    story = []
    for block in text.split("\n\n"):
        story.append(Paragraph(block.replace("\n", " "), styles["BodyText"]))
        story.append(Spacer(1, 8))
    SimpleDocTemplate(str(path), pagesize=A4).build(story)


def write_docx(path: Path, spec: dict) -> None:
    doc = Document()
    doc.add_heading(spec["title"], level=1)
    for paragraph in spec["paragraphs"]:
        doc.add_paragraph(paragraph)
    if spec["table"]:
        doc.add_heading("Order Details", level=2)
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for label, value in spec["table"]:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = label, value
    doc.save(str(path))


def main() -> None:
    DATA_DIR.mkdir(exist_ok=True)

    write_pdf(DATA_DIR / "complaint_001.pdf", COMPLAINT_001_PDF)
    (DATA_DIR / "complaint_002.txt").write_text(COMPLAINT_002_TXT, encoding="utf-8")
    write_docx(DATA_DIR / "complaint_003.docx", COMPLAINT_003_DOCX)
    write_pdf(DATA_DIR / "complaint_004.pdf", COMPLAINT_004_PDF)
    (DATA_DIR / "complaint_005.txt").write_text(COMPLAINT_005_TXT, encoding="utf-8")
    write_docx(DATA_DIR / "complaint_006.docx", COMPLAINT_006_DOCX)

    # Error-handling samples
    (DATA_DIR / "complaint_007_corrupt.pdf").write_bytes(b"%PDF-1.4\nthis file is truncated and corrupt")
    (DATA_DIR / "complaint_008.xlsx").write_bytes(b"not a supported format")

    print(f"Sample data written to {DATA_DIR}")


if __name__ == "__main__":
    main()
