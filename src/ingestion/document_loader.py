"""Document ingestion: discover files in the data folder and extract their text.

Supported formats: .txt, .pdf, .docx
"""

from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from src.logger import get_logger

logger = get_logger(__name__)

# Minimum characters of text for a document to be worth sending to the LLM.
MIN_TEXT_LENGTH = 30


class DocumentLoadError(Exception):
    """Raised when a file cannot be read or contains no usable text."""


@dataclass
class LoadedDocument:
    file_name: str
    file_path: Path
    file_type: str
    text: str

    @property
    def document_id(self) -> str:
        """Stable identifier used for output file names (e.g. complaint_001)."""
        return self.file_path.stem


def _read_txt(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentLoadError(f"Could not decode text file: {path.name}")


def _read_pdf(path: Path) -> str:
    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise DocumentLoadError(f"PDF is password protected: {path.name}")
        pages = [page.extract_text() or "" for page in reader.pages]
    except PdfReadError as exc:
        raise DocumentLoadError(f"Corrupt or unreadable PDF '{path.name}': {exc}") from exc
    return "\n".join(pages)


def _read_docx(path: Path) -> str:
    try:
        doc = DocxDocument(str(path))
    except Exception as exc:  # python-docx raises several unrelated exception types
        raise DocumentLoadError(f"Corrupt or unreadable DOCX '{path.name}': {exc}") from exc

    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


READERS = {
    ".txt": _read_txt,
    ".pdf": _read_pdf,
    ".docx": _read_docx,
}

SUPPORTED_EXTENSIONS = tuple(READERS)


def discover_documents(data_dir: Path) -> tuple[list[Path], list[Path]]:
    """Return (supported_files, skipped_files) found directly inside data_dir."""
    if not data_dir.is_dir():
        raise DocumentLoadError(f"Data directory not found: {data_dir}")

    supported, skipped = [], []
    for path in sorted(data_dir.iterdir()):
        if not path.is_file() or path.name.startswith("."):
            continue
        if path.suffix.lower() in READERS:
            supported.append(path)
        else:
            skipped.append(path)
    return supported, skipped


def load_document(path: Path) -> LoadedDocument:
    """Read a single file and return its normalised text content."""
    reader = READERS.get(path.suffix.lower())
    if reader is None:
        raise DocumentLoadError(f"Unsupported file type: {path.suffix}")

    try:
        raw_text = reader(path)
    except DocumentLoadError:
        raise
    except OSError as exc:
        raise DocumentLoadError(f"Could not open '{path.name}': {exc}") from exc

    text = "\n".join(line.strip() for line in raw_text.splitlines()).strip()
    if len(text) < MIN_TEXT_LENGTH:
        raise DocumentLoadError(
            f"'{path.name}' contains no usable text ({len(text)} characters extracted)"
        )

    logger.debug("Loaded %s (%d characters)", path.name, len(text))
    return LoadedDocument(
        file_name=path.name,
        file_path=path,
        file_type=path.suffix.lower().lstrip("."),
        text=text,
    )
